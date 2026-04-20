"""Helpers shared across reports."""
from openpyxl.utils import get_column_letter
from ..styles import (
    FONT_TITLE, FONT_SUBTITLE, FONT_HEADER, FONT_HEADER_WHITE, FONT_BODY,
    FONT_BODY_BOLD, FONT_BREAK, FILL_TITLE, FILL_HEADER_GRAY, FILL_GREEN_SOFT,
    FILL_DAY_BANNER, FILL_SECTION_HDR, CENTER, LEFT, BORDER, FILL_WORKING,
    FILL_HOLD, FILL_TRAVEL, FILL_OFF,
)


def write_title_block(ws, title: str, subtitle: str = "", n_cols: int = 10):
    """Cell A1: title row (merged across n_cols). A2: subtitle."""
    ws.cell(1, 1).value = title
    ws.cell(1, 1).font = FONT_TITLE
    ws.cell(1, 1).fill = FILL_TITLE
    ws.cell(1, 1).alignment = CENTER
    ws.merge_cells(start_row=1, end_row=1, start_column=1, end_column=n_cols)
    ws.row_dimensions[1].height = 26

    if subtitle:
        ws.cell(2, 1).value = subtitle
        ws.cell(2, 1).font = FONT_SUBTITLE
        ws.cell(2, 1).fill = FILL_TITLE
        ws.cell(2, 1).alignment = CENTER
        ws.merge_cells(start_row=2, end_row=2, start_column=1, end_column=n_cols)
        ws.row_dimensions[2].height = 18


def apply_print_settings(ws, fit_width=1, fit_height=0, orientation="landscape"):
    """Make the sheet export/print in a single readable page-width."""
    ws.page_setup.orientation = orientation
    ws.page_setup.fitToWidth = fit_width
    ws.page_setup.fitToHeight = fit_height
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4


def autosize(ws, min_col=1, max_col=None, max_width=60):
    if max_col is None:
        max_col = ws.max_column
    for c in range(min_col, max_col + 1):
        letter = get_column_letter(c)
        max_len = 0
        for r in range(1, ws.max_row + 1):
            v = ws.cell(r, c).value
            if v is None:
                continue
            L = max(len(line) for line in str(v).split("\n")) if str(v) else 0
            if L > max_len:
                max_len = L
        ws.column_dimensions[letter].width = min(max(12, max_len + 2), max_width)
