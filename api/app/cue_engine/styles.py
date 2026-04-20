"""
Shared cell styling for Cue v2 reports.
Matches the visual conventions of the Katakirr master-breakdown template.
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# -- fonts --
FONT_TITLE        = Font(name="Times New Roman", size=14, bold=True)
FONT_SUBTITLE     = Font(name="Times New Roman", size=10, bold=True, italic=True)
FONT_HEADER       = Font(name="Times New Roman", size=9,  bold=True)
FONT_HEADER_WHITE = Font(name="Times New Roman", size=9,  bold=True, color="FFFFFFFF")
FONT_BODY         = Font(name="Arial",            size=9)
FONT_BODY_BOLD    = Font(name="Arial",            size=9, bold=True)
FONT_DAY_BANNER   = Font(name="Times New Roman", size=12, bold=True)
FONT_BREAK        = Font(name="Arial",            size=9, bold=True)

# -- fills (matching Katakirr legend colors) --
FILL_TITLE       = PatternFill("solid", fgColor="FFF2CC")  # cream
FILL_HEADER_GRAY = PatternFill("solid", fgColor="CCCCCC")  # gray for DATE/DAY headers
FILL_CAMERA      = PatternFill("solid", fgColor="C00000")  # C1/C2/C3 (deep red)
FILL_LIGHT       = PatternFill("solid", fgColor="FF0000")  # L1/L2/L3 (red)
FILL_STEADICAM   = PatternFill("solid", fgColor="EC7C30")  # SL (orange)
FILL_PLAYBACK    = PatternFill("solid", fgColor="9BC2E6")  # PD (blue)
FILL_PT          = PatternFill("solid", fgColor="66FF99")  # PT (green)
FILL_SC          = PatternFill("solid", fgColor="ADAAAA")  # SC (gray)
FILL_JIB         = PatternFill("solid", fgColor="F4AF84")  # JJ (peach)
FILL_DRONE       = PatternFill("solid", fgColor="FF5050")  # DR (red-ish)
FILL_JR          = PatternFill("solid", fgColor="777777")  # JR (dark gray)
FILL_AD          = PatternFill("solid", fgColor="D5DCE3")  # AD (light gray)
FILL_PV          = PatternFill("solid", fgColor="A9D08E")  # PV (green)
FILL_CRANE       = PatternFill("solid", fgColor="FFC000")  # CR (yellow-orange)
FILL_PINK        = PatternFill("solid", fgColor="FF33CC")  # AH / TS
FILL_BLUEGRAY    = PatternFill("solid", fgColor="8496AF")  # AX / IC
FILL_LIGHTBLUE   = PatternFill("solid", fgColor="DDEBF7")  # HC / LL
FILL_DARKBLUE    = PatternFill("solid", fgColor="0070C0")  # VFX / SFX / GP
FILL_GREEN_SOFT  = PatternFill("solid", fgColor="D9EAD3")  # break day
FILL_CAST        = PatternFill("solid", fgColor="ADAAAA")  # cast columns
FILL_DAY_BANNER  = PatternFill("solid", fgColor="FFFF00")  # yellow (Nirmal Pathak style)
FILL_SECTION_HDR = PatternFill("solid", fgColor="00B0F0")  # cyan (Nirmal Pathak style)
FILL_WORKING     = PatternFill("solid", fgColor="A9D08E")  # green cell = working
FILL_HOLD        = PatternFill("solid", fgColor="FFF2CC")  # cream cell = hold
FILL_TRAVEL      = PatternFill("solid", fgColor="9BC2E6")  # blue cell = travel
FILL_OFF         = PatternFill("solid", fgColor="EEEEEE")  # light gray = off

# -- alignment --
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
TOPLEFT = Alignment(horizontal="left", vertical="top",  wrap_text=True)

# -- borders --
_thin = Side(style="thin", color="888888")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

# Column-code color map — used when rendering header rows in the breakdown.
UNIT_COLORS = {
    "C1": FILL_CAMERA, "C2": FILL_CAMERA, "C3": FILL_CAMERA,
    "L1": FILL_LIGHT,  "L2": FILL_LIGHT,  "L3": FILL_LIGHT,
    "SL": FILL_STEADICAM,
    "PD": FILL_PLAYBACK,
    "PT": FILL_PT,
    "SC": FILL_SC,
    "JJ": FILL_JIB,
    "DR": FILL_DRONE,
    "JR": FILL_JR,
    "AD": FILL_AD,
    "PV": FILL_PV,
    "CR": FILL_CRANE,
    "AH": FILL_PINK, "TS": FILL_PINK,
    "AX": FILL_BLUEGRAY, "IC": FILL_BLUEGRAY,
    "HC": FILL_LIGHTBLUE, "LL": FILL_LIGHTBLUE,
    "VFX": FILL_DARKBLUE, "SFX": FILL_DARKBLUE, "GP": FILL_DARKBLUE,
    "CRA": FILL_PT,
}

# Unit full names for legends / column headers in some reports
UNIT_NAMES = {
    "C1": "Camera 1", "C2": "Camera 2", "C3": "Camera 3",
    "L1": "Lights 1", "L2": "Lights 2", "L3": "Lights 3",
    "SL": "Steadicam",
    "PD": "Playback",
    "PT": "Props & Tech",
    "SC": "Stunt Co-ord",
    "JJ": "Jimmy Jib",
    "DR": "Drone",
    "JR": "Junior Artists",
    "AD": "Art Dept",
    "PV": "Picture Vehicle",
    "CR": "Crane",
    "AH": "Action / Hero",
    "AX": "Action Extras",
    "IC": "Intimacy Co-ord",
    "HC": "Hair / Costume",
    "VFX": "VFX",
    "LL": "Lighting Lab",
    "SFX": "SFX",
    "TS": "Tech Support",
    "CRA": "CRA",
    "GP": "Grip",
}
