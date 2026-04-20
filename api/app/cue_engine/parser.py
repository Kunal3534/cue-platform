"""
Master breakdown parser.

Accepts any xlsx that matches the Katakirr template:
  Row 1 : project title (any merged or single cell)
  Row 2 : shoot window description (optional)
  Row 3 : column headers — must include DATE, DAY, LOCATION, FILM LOCATION,
          Shift Timing, plus any unit codes (C1, L1, JR, PV, ...), plus
          CAST, SUPPORTING CAST, PICTURE VEHICLE, EXTRAS, NOTES.
  Row 4+: data rows — either a shoot day or a break/rehearsal day.

Optionally a second sheet "Cast" maps character → actor.

Optionally a third sheet "Scenes" maps shoot-days → scene-level detail:
  columns: SHOOT DAY, SC, DESCRIPTION, LOC, SUB-LOC, D/N, I/E, CHARACTERS,
           PC, TIME, CAM, SPL EQUIP, CODE.
"""
from __future__ import annotations
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Any
import re
import openpyxl


# ----------------------------- data classes -----------------------------

@dataclass
class Day:
    row_idx: int                          # source row (for debugging)
    date_raw: str                         # e.g. "15th March"
    day_label: str                        # e.g. "Shoot Day 1" or "Break Day"
    is_shoot: bool
    shoot_day_no: Optional[int] = None    # 1, 2, 3, ...
    location: str = ""                    # high-level, e.g. "GAGANBAWDA"
    film_location: str = ""               # specific, e.g. "ABHAY WADA"
    shift: str = ""                       # "6 AM to 6 PM"
    units: Dict[str, float] = field(default_factory=dict)  # C1 → 1.0, JR → 47, ...
    cast: List[str] = field(default_factory=list)
    supporting_cast: List[str] = field(default_factory=list)
    picture_vehicles: List[str] = field(default_factory=list)
    extras: str = ""
    notes: str = ""


@dataclass
class Project:
    title: str = "Untitled Production"
    window: str = ""
    unit_cols: List[str] = field(default_factory=list)  # in order found in template
    days: List[Day] = field(default_factory=list)
    cast_map: Dict[str, str] = field(default_factory=dict)  # character → actor
    scenes_by_day: Dict[str, List[Dict[str, Any]]] = field(default_factory=dict)
    # scenes_by_day["Shoot Day 1"] = [{"sc": "1-19A", "description": "...", ...}, ...]

    @property
    def shoot_days(self) -> List[Day]:
        return [d for d in self.days if d.is_shoot]


# ----------------------------- helpers -----------------------------

# Fixed-position header slots we always expect:
_META_HEADERS = {
    "DATE": "date",
    "DAY":  "day",
    "LOCATION": "location",
    "FILM LOCATION": "film_location",
    "SHIFT TIMING": "shift",
    "CAST": "cast",
    "SUPPORTING CAST": "supporting_cast",
    "PICTURE VEHICLE": "picture_vehicles",
    "EXTRAS": "extras",
    "NOTES": "notes",
}

def _norm(s: Any) -> str:
    if s is None:
        return ""
    return re.sub(r"\s+", " ", str(s).strip()).upper()


def _split_names(s: Any) -> List[str]:
    if s is None:
        return []
    txt = str(s).strip()
    if not txt:
        return []
    # split on comma / " / " / newline / semicolon
    parts = re.split(r"[,/;\n]+", txt)
    return [p.strip() for p in parts if p.strip()]


def _parse_shoot_day_no(label: str) -> Optional[int]:
    m = re.search(r"shoot\s*day\s*(\d+)", label, re.I)
    return int(m.group(1)) if m else None


def _find_header_row(ws) -> int:
    """Find the header row — the first row containing both 'DATE' and 'DAY'."""
    for r in range(1, min(ws.max_row, 20) + 1):
        vals = [_norm(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if "DATE" in vals and "DAY" in vals:
            return r
    raise ValueError("Could not locate header row (needs DATE and DAY).")


# ----------------------------- main parse -----------------------------

def load_breakdown(path: str) -> Project:
    wb = openpyxl.load_workbook(path, data_only=True)

    # primary sheet: first, or one named something like "Breakdown"
    breakdown_sheet = None
    for name in wb.sheetnames:
        if _norm(name) in ("BREAKDOWN", "SHEET1"):
            breakdown_sheet = wb[name]
            break
    if breakdown_sheet is None:
        breakdown_sheet = wb[wb.sheetnames[0]]

    proj = Project()

    # title from row 1 (any populated cell)
    for c in range(1, breakdown_sheet.max_column + 1):
        v = breakdown_sheet.cell(1, c).value
        if v:
            proj.title = str(v).strip()
            break
    # window from row 2
    for c in range(1, breakdown_sheet.max_column + 1):
        v = breakdown_sheet.cell(2, c).value
        if v:
            proj.window = str(v).strip()
            break

    header_row = _find_header_row(breakdown_sheet)
    # build column map: index → header label
    headers: Dict[int, str] = {}
    for c in range(1, breakdown_sheet.max_column + 1):
        label = _norm(breakdown_sheet.cell(header_row, c).value)
        if label:
            headers[c] = label

    # classify columns: meta vs unit
    meta_cols: Dict[str, int] = {}
    unit_cols: List[tuple] = []  # (col_idx, code)
    for c, label in headers.items():
        # match canonical meta names (case-insensitive)
        matched = False
        for canon, key in _META_HEADERS.items():
            if _norm(canon) == label:
                meta_cols[key] = c
                matched = True
                break
        if not matched:
            # anything else in row 3 that's not empty we treat as a unit code
            unit_cols.append((c, label))
    proj.unit_cols = [code for _, code in unit_cols]

    # iterate data rows
    for r in range(header_row + 1, breakdown_sheet.max_row + 1):
        date_val = breakdown_sheet.cell(r, meta_cols.get("date", 1)).value
        day_val  = breakdown_sheet.cell(r, meta_cols.get("day",  2)).value
        if not date_val and not day_val:
            continue
        day_label = str(day_val).strip() if day_val else ""
        date_raw = str(date_val).strip() if date_val else ""
        is_shoot = bool(_parse_shoot_day_no(day_label))

        day = Day(
            row_idx=r,
            date_raw=date_raw,
            day_label=day_label,
            is_shoot=is_shoot,
            shoot_day_no=_parse_shoot_day_no(day_label),
        )
        if is_shoot:
            day.location      = str(breakdown_sheet.cell(r, meta_cols.get("location", 3)).value or "").strip()
            day.film_location = str(breakdown_sheet.cell(r, meta_cols.get("film_location", 4)).value or "").strip()
            day.shift         = str(breakdown_sheet.cell(r, meta_cols.get("shift", 5)).value or "").strip()
            day.cast            = _split_names(breakdown_sheet.cell(r, meta_cols.get("cast", 32)).value)
            day.supporting_cast = _split_names(breakdown_sheet.cell(r, meta_cols.get("supporting_cast", 33)).value)
            day.picture_vehicles= _split_names(breakdown_sheet.cell(r, meta_cols.get("picture_vehicles", 34)).value)
            day.extras          = str(breakdown_sheet.cell(r, meta_cols.get("extras", 35)).value or "").strip()
            day.notes           = str(breakdown_sheet.cell(r, meta_cols.get("notes",  36)).value or "").strip()
            for col_idx, code in unit_cols:
                v = breakdown_sheet.cell(r, col_idx).value
                if isinstance(v, (int, float)) and v:
                    day.units[code] = float(v)
                elif v not in (None, ""):
                    try:
                        day.units[code] = float(v)
                    except (ValueError, TypeError):
                        pass
        proj.days.append(day)

    # ---- optional: cast mapping sheet ----
    cast_sheet = None
    for name in wb.sheetnames:
        n = _norm(name)
        if n in ("CAST", "CAST LIST", "ADVANCE PAYMENT SHEET", "ARTIST LIST"):
            cast_sheet = wb[name]
            break
    if cast_sheet is not None:
        # Look for columns resembling "Character Name" and "Name" (actor)
        for r in range(1, cast_sheet.max_row + 1):
            row_vals = [cast_sheet.cell(r, c).value for c in range(1, cast_sheet.max_column + 1)]
            if not any(row_vals):
                continue
            # detect header
            normed = [_norm(v) for v in row_vals]
            if "CHARACTER NAME" in normed or "CHARACTER" in normed:
                char_col = next(c for c, v in enumerate(normed) if v in ("CHARACTER NAME", "CHARACTER"))
                name_col = next((c for c, v in enumerate(normed) if v == "NAME"), char_col + 1)
                # subsequent rows
                for rr in range(r + 1, cast_sheet.max_row + 1):
                    ch = cast_sheet.cell(rr, char_col + 1).value
                    nm = cast_sheet.cell(rr, name_col + 1).value
                    if ch and nm:
                        proj.cast_map[str(ch).strip().upper()] = str(nm).strip()
                break

    # ---- optional: scenes sheet ----
    scenes_sheet = None
    for name in wb.sheetnames:
        if _norm(name) in ("SCENES", "SCENE BREAKDOWN"):
            scenes_sheet = wb[name]
            break
    if scenes_sheet is not None:
        # find header row
        hdr_r = None
        for r in range(1, min(scenes_sheet.max_row, 15) + 1):
            vals = [_norm(scenes_sheet.cell(r, c).value) for c in range(1, scenes_sheet.max_column + 1)]
            if "SC" in vals and ("DESCRIPTION" in vals or "DESC" in vals):
                hdr_r = r; break
        if hdr_r:
            headers2 = {c: _norm(scenes_sheet.cell(hdr_r, c).value)
                        for c in range(1, scenes_sheet.max_column + 1)}
            inv = {v: k for k, v in headers2.items() if v}
            def gv(r, key, alts=()):
                col = inv.get(_norm(key))
                for a in alts:
                    if col: break
                    col = inv.get(_norm(a))
                if col: return scenes_sheet.cell(r, col).value
                return None
            for r in range(hdr_r + 1, scenes_sheet.max_row + 1):
                shoot_day = gv(r, "SHOOT DAY", alts=("DAY",))
                sc = gv(r, "SC")
                if not shoot_day and not sc:
                    continue
                entry = {
                    "sc":          str(sc or "").strip(),
                    "description": str(gv(r, "DESCRIPTION", alts=("DESC",)) or "").strip(),
                    "loc":         str(gv(r, "LOC", alts=("LOCATION",)) or "").strip(),
                    "sub_loc":     str(gv(r, "SUB-LOC", alts=("SUBLOC",)) or "").strip(),
                    "dn":          str(gv(r, "D/N") or "").strip(),
                    "ie":          str(gv(r, "I/E") or "").strip(),
                    "characters":  str(gv(r, "CHARACTERS") or "").strip(),
                    "pc":          gv(r, "PC"),
                    "time":        str(gv(r, "TIME") or "").strip(),
                    "cam":         str(gv(r, "CAM A/B", alts=("CAM",)) or "").strip(),
                    "spl_equip":   str(gv(r, "SPL EQUIP") or "").strip(),
                    "code":        str(gv(r, "CODE") or "").strip(),
                }
                key = str(shoot_day or "").strip()
                proj.scenes_by_day.setdefault(key, []).append(entry)

    return proj
