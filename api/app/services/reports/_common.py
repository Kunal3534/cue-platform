"""Shared helpers for all v1 xlsx reports (accept ProjectGraph, not v0.1 objects)."""
from openpyxl.utils import get_column_letter
from ...cue_engine.styles import (  # reuse existing style constants
    FONT_TITLE, FONT_SUBTITLE, FONT_HEADER, FONT_HEADER_WHITE,
    FONT_BODY, FONT_BODY_BOLD, FONT_BREAK,
    FILL_TITLE, FILL_HEADER_GRAY, FILL_GREEN_SOFT,
    FILL_DAY_BANNER, FILL_SECTION_HDR,
    CENTER, LEFT, TOPLEFT, BORDER,
)


def write_title_block(ws, title: str, subtitle: str = "", n_cols: int = 10) -> None:
    ws.cell(1, 1).value = title
    ws.cell(1, 1).font = FONT_TITLE
    ws.cell(1, 1).fill = FILL_TITLE
    ws.cell(1, 1).alignment = CENTER
    ws.merge_cells(start_row=1, end_row=1, start_column=1, end_column=n_cols)
    ws.row_dimensions[1].height = 26

    if subtitle:
        ws.cell(2, 1).value = subtitle
        ws.cell(2, 1).font = FONT_SUBTITLE
        ws.cell(2, 1).fill = FILL_TITLE
        ws.cell(2, 1).alignment = CENTER
        ws.merge_cells(start_row=2, end_row=2, start_column=1, end_column=n_cols)
        ws.row_dimensions[2].height = 18


def write_header_row(ws, headers: list[str], row: int = 3) -> None:
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row, c)
        cell.value = h
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER_GRAY
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[row].height = 22


def apply_print_settings(ws, fit_width: int = 1, orientation: str = "landscape") -> None:
    ws.page_setup.orientation = orientation
    ws.page_setup.fitToWidth = fit_width
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4


def set_col_widths(ws, widths: dict[int, float]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def write_body_row(ws, row: int, values: list, bold_cols: set[int] | None = None) -> None:
    bold_cols = bold_cols or set()
    for c, v in enumerate(values, 1):
        cell = ws.cell(row, c)
        cell.value = v
        cell.font = FONT_BODY_BOLD if c in bold_cols else FONT_BODY
        cell.alignment = LEFT
        cell.border = BORDER
