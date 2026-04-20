"""14 — Requirements Breakdown xlsx report (all resources by shoot day)."""
from openpyxl import Workbook

from ._common import (
    write_title_block, apply_print_settings, set_col_widths,
    FONT_HEADER_WHITE, FONT_BODY, FONT_BODY_BOLD, FONT_BREAK,
    FILL_SECTION_HDR, FILL_GREEN_SOFT, FILL_DAY_BANNER,
    FONT_HEADER, CENTER, LEFT, BORDER,
)
from ..project_graph import ProjectGraph, fmt_pages, fmt_shift


def generate(pg: ProjectGraph, out_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Requirements"

    headers = ["RESOURCE TYPE", "DETAIL", "COUNT / SPEC", "STATUS / NOTES"]
    n = len(headers)
    write_title_block(ws, f"{pg.project.title} — REQUIREMENTS BREAKDOWN", subtitle=pg.project.window or "", n_cols=n)

    for c, h in enumerate(headers, 1):
        cell = ws.cell(3, c)
        cell.value = h
        cell.font = FONT_HEADER_WHITE
        cell.fill = FILL_SECTION_HDR
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[3].height = 22

    # Build per-day lookups
    animal_by_day: dict[str, list] = {}
    for animal in pg.animals:
        for did in animal.day_ids:
            animal_by_day.setdefault(did, []).append(animal)

    r = 4
    for day in pg.shoot_days:
        # Day banner
        sdn = day.shoot_day_number
        shift = fmt_shift(day.shift_start, day.shift_end)
        banner = f"DAY {sdn:02d} · {day.date.strftime('%-d %b %Y')} · {day.location_name or ''} · {shift}"
        cell = ws.cell(r, 1)
        cell.value = banner
        cell.font = FONT_HEADER
        cell.fill = FILL_DAY_BANNER
        cell.alignment = CENTER
        cell.border = BORDER
        ws.merge_cells(start_row=r, end_row=r, start_column=1, end_column=n)
        r += 1

        def _row(rtype, detail, count_spec, notes=""):
            for c, v in enumerate([rtype, detail, count_spec, notes], 1):
                cell = ws.cell(r, c)
                cell.value = v
                cell.font = FONT_BODY_BOLD if c == 1 else FONT_BODY
                cell.alignment = LEFT
                cell.border = BORDER

        # Camera
        for code, kind, cnt in day.camera_uses:
            _row("Camera/Equip", code, f"×{cnt}", kind.replace("_", " ").title())
            r += 1

        # JR
        if day.jr_count:
            _row("Junior Artists", "", day.jr_count, "")
            r += 1

        # Vehicles
        for vn in day.vehicle_names:
            _row("Picture Vehicle", vn, 1, "")
            r += 1

        # Animals
        for animal in animal_by_day.get(day.id, []):
            _row("Animal", animal.species, animal.count, animal.awbi_status or "")
            r += 1

    set_col_widths(ws, {1: 18, 2: 36, 3: 14, 4: 28})
    ws.freeze_panes = "A4"
    apply_print_settings(ws)
    wb.save(out_path)
