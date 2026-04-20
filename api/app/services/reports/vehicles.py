"""05 — Picture Vehicles xlsx report."""
from openpyxl import Workbook

from ._common import (
    write_title_block, apply_print_settings, set_col_widths,
    FONT_HEADER_WHITE, FONT_BODY, FONT_BODY_BOLD,
    FILL_SECTION_HDR, CENTER, LEFT, BORDER,
)
from ..project_graph import ProjectGraph


def generate(pg: ProjectGraph, out_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Picture Vehicles"

    headers = ["VEHICLE", "SPEC", "SOURCE", "VENDOR", "SHOOT DAYS", "NOTES"]
    n = len(headers)
    write_title_block(ws, f"{pg.project.title} — PICTURE VEHICLES", subtitle=pg.project.window or "", n_cols=n)

    for c, h in enumerate(headers, 1):
        cell = ws.cell(3, c)
        cell.value = h
        cell.font = FONT_HEADER_WHITE
        cell.fill = FILL_SECTION_HDR
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[3].height = 22

    day_id_to_sdn = {d.id: d.shoot_day_number for d in pg.shoot_days if d.shoot_day_number}

    r = 4
    for v in pg.vehicles:
        sdns = sorted(set(day_id_to_sdn[did] for did in v.day_ids if did in day_id_to_sdn))
        row_vals = [
            v.name,
            v.spec or "",
            v.source or "",
            v.source_vendor or "",
            ", ".join(str(s) for s in sdns),
            "",
        ]
        for c, val in enumerate(row_vals, 1):
            cell = ws.cell(r, c)
            cell.value = val
            cell.font = FONT_BODY_BOLD if c == 1 else FONT_BODY
            cell.alignment = LEFT
            cell.border = BORDER
        r += 1

    set_col_widths(ws, {1: 28, 2: 36, 3: 14, 4: 28, 5: 22, 6: 30})
    ws.freeze_panes = "A4"
    apply_print_settings(ws)
    wb.save(out_path)
