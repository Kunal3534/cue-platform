"""07 — Junior Artists xlsx report."""
from openpyxl import Workbook

from ._common import (
    write_title_block, apply_print_settings, set_col_widths,
    FONT_HEADER_WHITE, FONT_BODY, FONT_BODY_BOLD,
    FILL_SECTION_HDR, CENTER, LEFT, BORDER,
)
from ..project_graph import ProjectGraph


def generate(pg: ProjectGraph, out_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Junior Artists"

    headers = ["SHOOT DAY", "DATE", "GRADE", "COUNT", "VENDOR", "CONTEXT", "WARDROBE NOTES"]
    n = len(headers)
    write_title_block(ws, f"{pg.project.title} — JUNIOR ARTISTS", subtitle=pg.project.window or "", n_cols=n)

    for c, h in enumerate(headers, 1):
        cell = ws.cell(3, c)
        cell.value = h
        cell.font = FONT_HEADER_WHITE
        cell.fill = FILL_SECTION_HDR
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[3].height = 22

    r = 4
    for jr in pg.jr_calls:
        row_vals = [
            str(jr.shoot_day_number) if jr.shoot_day_number else "",
            jr.day_date.strftime("%-d %b %Y") if jr.day_date else "",
            jr.grade.title(),
            jr.count,
            jr.vendor_name or "",
            jr.context or "",
            "",
        ]
        for c, val in enumerate(row_vals, 1):
            cell = ws.cell(r, c)
            cell.value = val
            cell.font = FONT_BODY_BOLD if c in (1, 4) else FONT_BODY
            cell.alignment = CENTER if c in (1, 2, 3, 4) else LEFT
            cell.border = BORDER
        r += 1

    # Total row
    if pg.jr_calls:
        total = sum(j.count for j in pg.jr_calls)
        ws.cell(r, 1).value = "TOTAL"
        ws.cell(r, 4).value = total
        for c in range(1, n + 1):
            ws.cell(r, c).font = FONT_BODY_BOLD
            ws.cell(r, c).border = BORDER

    set_col_widths(ws, {1: 10, 2: 14, 3: 14, 4: 8, 5: 28, 6: 36, 7: 30})
    ws.freeze_panes = "A4"
    apply_print_settings(ws)
    wb.save(out_path)
