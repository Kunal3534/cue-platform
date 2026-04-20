"""01 — Main Schedule xlsx report (v1, accepts ProjectGraph)."""
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from ._common import (
    write_title_block, write_header_row, apply_print_settings, set_col_widths,
    FONT_HEADER, FONT_BODY, FONT_BODY_BOLD, FONT_BREAK,
    FILL_GREEN_SOFT, FILL_DAY_BANNER, FILL_SECTION_HDR, FILL_HEADER_GRAY,
    FONT_HEADER_WHITE, CENTER, LEFT, BORDER,
)
from ..project_graph import ProjectGraph, fmt_pages, fmt_shift


def generate(pg: ProjectGraph, out_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Main Schedule"

    headers = ["SC", "DESCRIPTION", "LOC", "D/N", "PAGES", "CHARACTERS", "SHIFT", "CAM UNITS", "JR", "NOTES"]
    n = len(headers)
    write_title_block(ws, f"{pg.project.title} — MAIN SCHEDULE", subtitle=pg.project.window or "", n_cols=n)

    r = 3
    for day in pg.days:
        if day.kind != "shoot":
            label = day.date.strftime("%-d %b %Y") + f" — {day.kind.upper()}"
            if day.banner_text:
                label += f" · {day.banner_text}"
            for c in range(1, n + 1):
                ws.cell(r, c).fill = FILL_GREEN_SOFT
                ws.cell(r, c).font = FONT_BREAK
                ws.cell(r, c).alignment = CENTER
                ws.cell(r, c).border = BORDER
            ws.cell(r, 1).value = label
            ws.merge_cells(start_row=r, end_row=r, start_column=1, end_column=n)
            r += 1
            continue

        # Day banner
        sdn = day.shoot_day_number
        loc = day.location_name or ""
        shift = fmt_shift(day.shift_start, day.shift_end)
        banner_val = f"DAY {sdn:02d} · {day.date.strftime('%-d %b')} · {loc} · {shift}"
        if day.banner_text:
            banner_val += f" · {day.banner_text}"
        cell = ws.cell(r, 1)
        cell.value = banner_val
        cell.font = FONT_HEADER
        cell.fill = FILL_DAY_BANNER
        cell.alignment = CENTER
        cell.border = BORDER
        ws.merge_cells(start_row=r, end_row=r, start_column=1, end_column=n)
        ws.row_dimensions[r].height = 20
        r += 1

        # Column headers under each day
        for c, h in enumerate(headers, 1):
            cell = ws.cell(r, c)
            cell.value = h
            cell.font = FONT_HEADER_WHITE
            cell.fill = FILL_SECTION_HDR
            cell.alignment = CENTER
            cell.border = BORDER
        r += 1

        cam_str = ", ".join(f"{code}×{cnt}" if cnt > 1 else code
                            for code, _kind, cnt in day.camera_uses)

        if not day.scenes:
            row_vals = [
                "—", day.location_name or "", "", "",
                "", ", ".join(day.cast), shift, cam_str,
                str(day.jr_count) if day.jr_count else "",
                day.notes or "",
            ]
            for c, v in enumerate(row_vals, 1):
                cell = ws.cell(r, c)
                cell.value = v
                cell.font = FONT_BODY
                cell.alignment = LEFT
                cell.border = BORDER
            r += 1
        else:
            for sc in day.scenes:
                row_vals = [
                    sc.scene_number,
                    sc.title or sc.description or "",
                    sc.location_name or day.location_name or "",
                    sc.dn.upper() if sc.dn else "",
                    fmt_pages(sc.page_eighths),
                    ", ".join(sc.characters),
                    shift,
                    cam_str,
                    str(day.jr_count) if day.jr_count else "",
                    day.notes or "",
                ]
                for c, v in enumerate(row_vals, 1):
                    cell = ws.cell(r, c)
                    cell.value = v
                    cell.font = FONT_BODY_BOLD if c == 1 else FONT_BODY
                    cell.alignment = LEFT
                    cell.border = BORDER
                r += 1

    set_col_widths(ws, {1: 8, 2: 48, 3: 20, 4: 8, 5: 8, 6: 36, 7: 14, 8: 20, 9: 6, 10: 36})
    ws.freeze_panes = "A4"
    apply_print_settings(ws)
    wb.save(out_path)
