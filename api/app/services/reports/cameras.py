"""06 — Camera Units xlsx report."""
from openpyxl import Workbook

from ._common import (
    write_title_block, apply_print_settings, set_col_widths,
    FONT_HEADER_WHITE, FONT_BODY, FONT_BODY_BOLD,
    FILL_SECTION_HDR, CENTER, LEFT, BORDER,
)
from ..project_graph import ProjectGraph


def generate(pg: ProjectGraph, out_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Camera Units"

    headers = ["CODE", "KIND", "DESCRIPTION", "RENTAL VENDOR", "OPERATOR", "SHOOT DAYS"]
    n = len(headers)
    write_title_block(ws, f"{pg.project.title} — CAMERA UNITS", subtitle=pg.project.window or "", n_cols=n)

    for c, h in enumerate(headers, 1):
        cell = ws.cell(3, c)
        cell.value = h
        cell.font = FONT_HEADER_WHITE
        cell.fill = FILL_SECTION_HDR
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[3].height = 22

    # Build unit_code → shoot day numbers
    unit_sdns: dict[str, list[int]] = {}
    for day in pg.shoot_days:
        if day.shoot_day_number is None:
            continue
        for code, _kind, _cnt in day.camera_uses:
            unit_sdns.setdefault(code, []).append(day.shoot_day_number)

    r = 4
    for cu in pg.camera_units:
        sdns = sorted(set(unit_sdns.get(cu.unit_code, [])))
        row_vals = [
            cu.unit_code,
            cu.kind.replace("_", " ").title() if cu.kind else "",
            cu.description or "",
            cu.rental_vendor or "",
            cu.operator_name or "",
            ", ".join(str(s) for s in sdns),
        ]
        for c, val in enumerate(row_vals, 1):
            cell = ws.cell(r, c)
            cell.value = val
            cell.font = FONT_BODY_BOLD if c == 1 else FONT_BODY
            cell.alignment = LEFT
            cell.border = BORDER
        r += 1

    set_col_widths(ws, {1: 8, 2: 14, 3: 36, 4: 28, 5: 28, 6: 24})
    ws.freeze_panes = "A4"
    apply_print_settings(ws)
    wb.save(out_path)
