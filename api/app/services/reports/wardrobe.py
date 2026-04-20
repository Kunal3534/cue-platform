"""13 — Wardrobe Looks xlsx report."""
from openpyxl import Workbook

from ._common import (
    write_title_block, apply_print_settings, set_col_widths,
    FONT_HEADER_WHITE, FONT_BODY, FONT_BODY_BOLD,
    FILL_SECTION_HDR, CENTER, LEFT, BORDER,
)
from ..project_graph import ProjectGraph


def generate(pg: ProjectGraph, out_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Wardrobe"

    headers = ["CHARACTER", "LOOK", "ERA", "DESCRIPTION", "SOURCE", "VENDOR", "MULTIPLES"]
    n = len(headers)
    write_title_block(ws, f"{pg.project.title} — WARDROBE", subtitle=pg.project.window or "", n_cols=n)

    for c, h in enumerate(headers, 1):
        cell = ws.cell(3, c)
        cell.value = h
        cell.font = FONT_HEADER_WHITE
        cell.fill = FILL_SECTION_HDR
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[3].height = 22

    r = 4
    for look in pg.wardrobe_looks:
        row_vals = [
            look.character_name.title(),
            look.look_id or "",
            look.era.replace("_", " ").title() if look.era else "",
            look.description or "",
            look.source.replace("_", " ").title() if look.source else "",
            look.source_vendor or "",
            look.multiples_count,
        ]
        for c, val in enumerate(row_vals, 1):
            cell = ws.cell(r, c)
            cell.value = val
            cell.font = FONT_BODY_BOLD if c == 1 else FONT_BODY
            cell.alignment = LEFT
            cell.border = BORDER
        r += 1

    set_col_widths(ws, {1: 24, 2: 8, 3: 14, 4: 36, 5: 18, 6: 28, 7: 10})
    ws.freeze_panes = "A4"
    apply_print_settings(ws)
    wb.save(out_path)
