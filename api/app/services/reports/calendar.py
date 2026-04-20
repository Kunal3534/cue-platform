"""04 — Calendar xlsx report."""
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from ._common import (
    write_title_block, apply_print_settings, set_col_widths,
    FONT_HEADER, FONT_HEADER_WHITE, FONT_BODY, FONT_BODY_BOLD, FONT_BREAK,
    FILL_SECTION_HDR, FILL_GREEN_SOFT, CENTER, LEFT, BORDER,
)
from ..project_graph import ProjectGraph, fmt_shift

_DOW = ["MON", "TUE", "WED", "THU", "FRI", "SAT", "SUN"]


def generate(pg: ProjectGraph, out_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Calendar"

    headers = ["DATE", "DOW", "KIND", "SHOOT DAY", "LOCATION", "SHIFT", "CAST (PRINCIPAL)", "JR", "NOTES"]
    n = len(headers)
    write_title_block(ws, f"{pg.project.title} — CALENDAR", subtitle=pg.project.window or "", n_cols=n)

    for c, h in enumerate(headers, 1):
        cell = ws.cell(3, c)
        cell.value = h
        cell.font = FONT_HEADER_WHITE
        cell.fill = FILL_SECTION_HDR
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[3].height = 22

    r = 4
    for day in pg.days:
        is_shoot = day.kind == "shoot"
        principals = [c for c in day.cast if any(
            ch.canonical_name == c and ch.is_principal for ch in pg.characters
        )]
        row_vals = [
            day.date.strftime("%-d %b %Y"),
            _DOW[day.date.weekday()],
            day.kind.upper(),
            str(day.shoot_day_number) if day.shoot_day_number else "",
            day.location_name or "",
            fmt_shift(day.shift_start, day.shift_end),
            ", ".join(principals) if is_shoot else "",
            str(day.jr_count) if day.jr_count else "",
            day.notes or "",
        ]
        for c, v in enumerate(row_vals, 1):
            cell = ws.cell(r, c)
            cell.value = v
            cell.font = FONT_BODY_BOLD if is_shoot and c == 1 else FONT_BODY
            cell.alignment = CENTER if c in (1, 2, 3, 4, 6, 8) else LEFT
            cell.border = BORDER
            if not is_shoot:
                cell.fill = FILL_GREEN_SOFT
        r += 1

    set_col_widths(ws, {1: 14, 2: 6, 3: 10, 4: 10, 5: 24, 6: 14, 7: 40, 8: 6, 9: 30})
    ws.freeze_panes = "A4"
    apply_print_settings(ws)
    wb.save(out_path)
