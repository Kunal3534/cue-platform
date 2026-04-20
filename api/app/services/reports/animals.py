"""09 — Animals xlsx report."""
from openpyxl import Workbook

from ._common import (
    write_title_block, apply_print_settings, set_col_widths,
    FONT_HEADER_WHITE, FONT_BODY, FONT_BODY_BOLD,
    FILL_SECTION_HDR, CENTER, LEFT, BORDER,
)
from ..project_graph import ProjectGraph


def generate(pg: ProjectGraph, out_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Animals"

    headers = ["SPECIES", "COUNT", "HANDLER", "VET ON SET", "AWBI STATUS", "SHOOT DAYS"]
    n = len(headers)
    write_title_block(ws, f"{pg.project.title} — ANIMALS", subtitle=pg.project.window or "", n_cols=n)

    for c, h in enumerate(headers, 1):
        cell = ws.cell(3, c)
        cell.value = h
        cell.font = FONT_HEADER_WHITE
        cell.fill = FILL_SECTION_HDR
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[3].height = 22

    day_id_to_sdn = {d.id: d.shoot_day_number for d in pg.shoot_days if d.shoot_day_number}

    r = 4
    for animal in pg.animals:
        sdns = sorted(set(day_id_to_sdn[did] for did in animal.day_ids if did in day_id_to_sdn))
        row_vals = [
            animal.species,
            animal.count,
            animal.handler_name or "",
            "Yes" if animal.vet_on_set else "No",
            animal.awbi_status.replace("_", " ").title() if animal.awbi_status else "",
            ", ".join(str(s) for s in sdns),
        ]
        for c, val in enumerate(row_vals, 1):
            cell = ws.cell(r, c)
            cell.value = val
            cell.font = FONT_BODY_BOLD if c == 1 else FONT_BODY
            cell.alignment = LEFT
            cell.border = BORDER
        r += 1

    set_col_widths(ws, {1: 28, 2: 8, 3: 28, 4: 12, 5: 18, 6: 20})
    ws.freeze_panes = "A4"
    apply_print_settings(ws)
    wb.save(out_path)
