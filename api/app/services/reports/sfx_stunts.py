"""11 — SFX & Stunts xlsx report."""
from openpyxl import Workbook

from ._common import (
    write_title_block, apply_print_settings, set_col_widths,
    FONT_HEADER_WHITE, FONT_BODY, FONT_BODY_BOLD,
    FILL_SECTION_HDR, CENTER, LEFT, BORDER,
)
from ..project_graph import ProjectGraph


def generate(pg: ProjectGraph, out_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "SFX & Stunts"

    headers = ["EFFECT KIND", "SCENE", "COORDINATOR", "MEDIC", "PERMITS", "STATUS"]
    n = len(headers)
    write_title_block(ws, f"{pg.project.title} — SFX & STUNTS", subtitle=pg.project.window or "", n_cols=n)

    for c, h in enumerate(headers, 1):
        cell = ws.cell(3, c)
        cell.value = h
        cell.font = FONT_HEADER_WHITE
        cell.fill = FILL_SECTION_HDR
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[3].height = 22

    r = 4
    for effect in pg.sfx_effects:
        permits = effect.permits_required
        if isinstance(permits, list):
            permits_str = ", ".join(str(p) for p in permits)
        elif permits:
            permits_str = str(permits)
        else:
            permits_str = ""

        row_vals = [
            effect.effect_kind.replace("_", " ").title() if effect.effect_kind else "",
            effect.scene_number or "",
            effect.coordinator or "",
            "Yes" if effect.medic_required else "No",
            permits_str,
            effect.status.replace("_", " ").title() if effect.status else "",
        ]
        for c, val in enumerate(row_vals, 1):
            cell = ws.cell(r, c)
            cell.value = val
            cell.font = FONT_BODY_BOLD if c == 1 else FONT_BODY
            cell.alignment = LEFT
            cell.border = BORDER
        r += 1

    set_col_widths(ws, {1: 24, 2: 10, 3: 28, 4: 8, 5: 28, 6: 16})
    ws.freeze_panes = "A4"
    apply_print_settings(ws)
    wb.save(out_path)
