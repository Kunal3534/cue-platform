"""10 — VFX Shots xlsx report."""
from openpyxl import Workbook

from ._common import (
    write_title_block, apply_print_settings, set_col_widths,
    FONT_HEADER_WHITE, FONT_BODY, FONT_BODY_BOLD,
    FILL_SECTION_HDR, CENTER, LEFT, BORDER,
)
from ..project_graph import ProjectGraph


def generate(pg: ProjectGraph, out_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "VFX"

    headers = ["SHOT#", "SCENE", "TYPE", "COMPLEXITY", "PLATES", "VENDOR", "BID STATUS", "STATUS"]
    n = len(headers)
    write_title_block(ws, f"{pg.project.title} — VFX SHOTS", subtitle=pg.project.window or "", n_cols=n)

    for c, h in enumerate(headers, 1):
        cell = ws.cell(3, c)
        cell.value = h
        cell.font = FONT_HEADER_WHITE
        cell.fill = FILL_SECTION_HDR
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[3].height = 22

    r = 4
    for shot in pg.vfx_shots:
        row_vals = [
            shot.shot_number,
            shot.scene_number or "",
            shot.shot_type.replace("_", " ").title() if shot.shot_type else "",
            shot.complexity.title() if shot.complexity else "",
            shot.plates_required,
            shot.vendor or "",
            shot.bid_status.replace("_", " ").title() if shot.bid_status else "",
            shot.status.replace("_", " ").title() if shot.status else "",
        ]
        for c, val in enumerate(row_vals, 1):
            cell = ws.cell(r, c)
            cell.value = val
            cell.font = FONT_BODY_BOLD if c == 1 else FONT_BODY
            cell.alignment = LEFT
            cell.border = BORDER
        r += 1

    set_col_widths(ws, {1: 10, 2: 10, 3: 22, 4: 12, 5: 8, 6: 28, 7: 16, 8: 16})
    ws.freeze_panes = "A4"
    apply_print_settings(ws)
    wb.save(out_path)
