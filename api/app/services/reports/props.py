"""08 — Props xlsx report."""
from openpyxl import Workbook

from ._common import (
    write_title_block, apply_print_settings, set_col_widths,
    FONT_HEADER_WHITE, FONT_BODY, FONT_BODY_BOLD,
    FILL_SECTION_HDR, CENTER, LEFT, BORDER,
)
from ..project_graph import ProjectGraph


def generate(pg: ProjectGraph, out_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Props"

    headers = ["PROP", "CATEGORY", "SOURCE", "VENDOR", "LEAD DAYS", "SCENES", "STATUS"]
    n = len(headers)
    write_title_block(ws, f"{pg.project.title} — PROPS", subtitle=pg.project.window or "", n_cols=n)

    for c, h in enumerate(headers, 1):
        cell = ws.cell(3, c)
        cell.value = h
        cell.font = FONT_HEADER_WHITE
        cell.fill = FILL_SECTION_HDR
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[3].height = 22

    # Build prop_name → scene numbers
    prop_scenes: dict[str, list[str]] = {}
    for link in pg.scene_props:
        prop_scenes.setdefault(link.prop_name, []).append(link.scene_number)

    r = 4
    for prop in pg.props:
        scenes_str = ", ".join(sorted(set(prop_scenes.get(prop.name, []))))
        row_vals = [
            prop.name,
            prop.category.replace("_", " ").title() if prop.category else "",
            prop.source.replace("_", " ").title() if prop.source else "",
            prop.source_vendor or "",
            str(prop.lead_days) if prop.lead_days else "",
            scenes_str,
            prop.status.replace("_", " ").title() if prop.status else "",
        ]
        for c, val in enumerate(row_vals, 1):
            cell = ws.cell(r, c)
            cell.value = val
            cell.font = FONT_BODY_BOLD if c == 1 else FONT_BODY
            cell.alignment = LEFT
            cell.border = BORDER
        r += 1

    set_col_widths(ws, {1: 32, 2: 18, 3: 16, 4: 28, 5: 10, 6: 24, 7: 16})
    ws.freeze_panes = "A4"
    apply_print_settings(ws)
    wb.save(out_path)
