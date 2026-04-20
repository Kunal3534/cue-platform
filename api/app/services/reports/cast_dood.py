"""02 — Cast DOOD (Day-Out-Of-Days) xlsx report."""
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from ._common import (
    write_title_block, apply_print_settings, set_col_widths,
    FONT_HEADER, FONT_HEADER_WHITE, FONT_BODY, FONT_BODY_BOLD,
    FILL_SECTION_HDR, CENTER, LEFT, BORDER,
)
from ...cue_engine.styles import FILL_WORKING, FILL_HOLD, FILL_OFF
from ..project_graph import ProjectGraph

_FILL_W = FILL_WORKING
_FILL_H = FILL_HOLD
_FILL_DOT = FILL_OFF


def generate(pg: ProjectGraph, out_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cast DOOD"

    shoot_days = pg.shoot_days
    sdns = [d.shoot_day_number for d in shoot_days if d.shoot_day_number is not None]
    n_fixed = 3  # Character, Actor, P/NP
    n_cols = n_fixed + len(sdns) + 3  # + W, H, T summary cols

    write_title_block(ws, f"{pg.project.title} — CAST DOOD", subtitle=pg.project.window or "", n_cols=n_cols)

    # Header row
    r = 3
    headers_fixed = ["CHARACTER", "ACTOR", "P/NP"]
    headers_days = [str(sdn) for sdn in sdns]
    headers_summary = ["W", "H", "T"]
    all_headers = headers_fixed + headers_days + headers_summary

    for c, h in enumerate(all_headers, 1):
        cell = ws.cell(r, c)
        cell.value = h
        cell.font = FONT_HEADER_WHITE
        cell.fill = FILL_SECTION_HDR
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[r].height = 22
    r += 1

    # Date sub-header
    for i, day in enumerate(shoot_days):
        if day.shoot_day_number is None:
            continue
        c = n_fixed + i + 1
        cell = ws.cell(r, c)
        cell.value = day.date.strftime("%-d %b")
        cell.font = FONT_BODY
        cell.alignment = CENTER
        cell.border = BORDER
    r += 1

    # Character rows
    char_by_id = {c.id: c for c in pg.characters}
    sdn_to_idx = {sdn: i for i, sdn in enumerate(sdns)}

    for char in pg.characters:
        dood = pg.character_dood.get(char.id, {})
        w_count = sum(1 for v in dood.values() if v == "W")
        h_count = sum(1 for v in dood.values() if v == "H")
        t_count = w_count + h_count

        row_vals = [char.display_name, char.actor_name or "", "P" if char.is_principal else "NP"]
        for sdn in sdns:
            row_vals.append(dood.get(sdn, "."))
        row_vals += [str(w_count), str(h_count), str(t_count)]

        for c, v in enumerate(row_vals, 1):
            cell = ws.cell(r, c)
            cell.value = v
            cell.font = FONT_BODY_BOLD if c <= n_fixed else FONT_BODY
            cell.alignment = CENTER if c > n_fixed else LEFT
            cell.border = BORDER
            # DOOD cell fill
            if c > n_fixed and c <= n_fixed + len(sdns):
                if v == "W":
                    cell.fill = _FILL_W
                elif v == "H":
                    cell.fill = _FILL_H
                else:
                    cell.fill = _FILL_DOT
        r += 1

    # Day totals footer
    ws.cell(r, 1).value = "Day total (W)"
    ws.cell(r, 1).font = FONT_BODY_BOLD
    ws.cell(r, 1).border = BORDER
    ws.merge_cells(start_row=r, end_row=r, start_column=1, end_column=n_fixed)
    for i, sdn in enumerate(sdns):
        total = sum(1 for ch in pg.characters if pg.character_dood.get(ch.id, {}).get(sdn) == "W")
        c = n_fixed + i + 1
        cell = ws.cell(r, c)
        cell.value = total
        cell.font = FONT_BODY_BOLD
        cell.alignment = CENTER
        cell.border = BORDER

    set_col_widths(ws, {1: 22, 2: 24, 3: 6})
    for i in range(len(sdns)):
        ws.column_dimensions[ws.cell(3, n_fixed + i + 1).column_letter].width = 5
    ws.freeze_panes = "D4"
    apply_print_settings(ws)
    wb.save(out_path)
