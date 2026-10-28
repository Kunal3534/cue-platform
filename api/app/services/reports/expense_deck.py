"""Expense report as a slide deck (.pptx).

The producers hand this to the Condé Nast accounts team as Google Slides —
a .pptx uploaded to Google Drive opens natively in Slides, so we build one:
cover → summary by shoot → line items per shoot → receipts appendix.

Styled to match the GQ MOTY calendar: paper ground, ink text, hairline rules.
"""
from __future__ import annotations

import io
from datetime import date as date_cls, datetime

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

PAPER = RGBColor(0xEF, 0xED, 0xE8)
INK = RGBColor(0x14, 0x14, 0x18)
MUTED = RGBColor(0x8A, 0x87, 0x80)
GOLD = RGBColor(0xB0, 0x8D, 0x2E)
HAIR = RGBColor(0xD8, 0xD5, 0xCE)

SERIF = "Georgia"
SANS = "Helvetica Neue"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

CATEGORY_LABELS = {
    "travel": "Travel", "food": "Food & catering", "location": "Location",
    "equipment": "Equipment", "styling": "Styling", "hmu": "Hair & makeup",
    "production": "Production", "misc": "Miscellaneous",
}


def inr(n: float) -> str:
    """Indian-style grouping: 12,34,567.89"""
    neg = n < 0
    n = abs(n)
    whole = int(n)
    frac = round((n - whole) * 100)
    if frac == 100:  # float edge, e.g. 9.999
        whole, frac = whole + 1, 0
    s = str(whole)
    if len(s) > 3:
        head, tail = s[:-3], s[-3:]
        parts = []
        while len(head) > 2:
            parts.insert(0, head[-2:])
            head = head[:-2]
        if head:
            parts.insert(0, head)
        s = ",".join(parts + [tail])
    return f"{'−' if neg else ''}₹{s}.{frac:02d}"


def _fmt_day(iso: str) -> str:
    d = date_cls.fromisoformat(iso)
    return d.strftime("%A, ") + f"{d.day} " + d.strftime("%B %Y")


def _blank(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = PAPER
    bg.line.fill.background()
    bg.shadow.inherit = False
    return slide


def _text(slide, x, y, w, h, text, *, size, font=SANS, color=INK, bold=False,
          align=PP_ALIGN.LEFT, spacing=0.0, upper=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text.upper() if upper else text
    f = run.font
    f.name = font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    if spacing:
        # letterspacing isn't exposed by python-pptx's API; fake it lightly
        run.text = (" ".join(run.text) if spacing >= 2 else run.text)
    return box


def _rule(slide, x, y, w, color=HAIR, height=Pt(1)):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Emu(int(height)))
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    bar.shadow.inherit = False
    return bar


def _style_cell(cell, *, size=11, bold=False, color=INK, align=PP_ALIGN.LEFT,
                fill=None):
    if fill is None:
        cell.fill.background()
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    cell.margin_left = Inches(0.08)
    cell.margin_right = Inches(0.08)
    cell.margin_top = Inches(0.04)
    cell.margin_bottom = Inches(0.04)
    for p in cell.text_frame.paragraphs:
        p.alignment = align
        for run in p.runs or [p.add_run()]:
            run.font.name = SANS
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color


def _header(slide, title: str, sub: str | None = None):
    _text(slide, Inches(0.7), Inches(0.45), Inches(8), Inches(0.4),
          "GQ MEN OF THE YEAR", size=11, color=MUTED, spacing=2, upper=True)
    _text(slide, Inches(0.7), Inches(0.85), Inches(12), Inches(0.8),
          title, size=30, font=SERIF, color=INK)
    if sub:
        _text(slide, Inches(0.7), Inches(1.55), Inches(12), Inches(0.4),
              sub, size=12, color=MUTED)
    _rule(slide, Inches(0.7), Inches(2.05), SLIDE_W - Inches(1.4))


def _receipt_jpeg(data: bytes, content_type: str | None) -> bytes | None:
    """Receipt file → a deck-sized JPEG (PDFs: first page)."""
    try:
        from PIL import Image
        if (content_type or "").endswith("pdf") or data[:5] == b"%PDF-":
            import pypdfium2 as pdfium
            doc = pdfium.PdfDocument(data)
            img = doc[0].render(scale=2.0).to_pil()
            doc.close()
        else:
            img = Image.open(io.BytesIO(data))
        img = img.convert("RGB")
        img.thumbnail((1400, 1400))
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=80)
        return buf.getvalue()
    except Exception:
        return None


def build_deck(
    project_title: str,
    expenses: list[dict],
    shoot_headlines: dict[str, str | None],
) -> bytes:
    from ...services import storage

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    grand = sum(e["amount"] or 0 for e in expenses)

    # Group by shoot day (None → general bucket at the end)
    groups: dict[str | None, list[dict]] = {}
    for e in expenses:
        groups.setdefault(e["date"], []).append(e)
    dated = sorted((k for k in groups if k), key=str)
    ordered: list[str | None] = [*dated, *([None] if None in groups else [])]

    def group_label(key: str | None) -> str:
        if key is None:
            return "General (not tied to one shoot)"
        head = shoot_headlines.get(key)
        return f"{_fmt_day(key)}" + (f" — {head}" if head else "")

    # ── Cover ──────────────────────────────────────────────────────────
    slide = _blank(prs)
    _text(slide, Inches(0.7), Inches(2.2), Inches(12), Inches(0.5),
          "GQ MEN OF THE YEAR", size=14, color=GOLD, spacing=2, upper=True)
    _text(slide, Inches(0.7), Inches(2.75), Inches(12), Inches(1.4),
          "Expense Report", size=60, font=SERIF)
    _text(slide, Inches(0.7), Inches(4.25), Inches(12), Inches(0.5),
          project_title, size=16, color=MUTED)
    _rule(slide, Inches(0.7), Inches(5.0), Inches(3.4), color=INK, height=Pt(2))
    _text(slide, Inches(0.7), Inches(5.25), Inches(12), Inches(0.6),
          f"Total spend  {inr(grand)}", size=20, bold=True)
    _text(slide, Inches(0.7), Inches(6.6), Inches(12), Inches(0.4),
          f"Prepared {datetime.now().strftime('%d %B %Y')} · Produced on Cue",
          size=10, color=MUTED)

    # ── Summary by shoot ───────────────────────────────────────────────
    slide = _blank(prs)
    _header(slide, "Summary by shoot")
    n_rows = len(ordered) + 2  # header + total
    table = slide.shapes.add_table(
        n_rows, 3, Inches(0.7), Inches(2.35), SLIDE_W - Inches(1.4),
        Inches(0.42 * n_rows),
    ).table
    table.columns[0].width = Inches(6.6)
    table.columns[1].width = Inches(2.4)
    table.columns[2].width = Inches(2.93)
    for c, label in enumerate(("Shoot", "Items", "Amount")):
        cell = table.cell(0, c)
        cell.text = label.upper()
        _style_cell(cell, size=10, bold=True, color=MUTED,
                    align=PP_ALIGN.RIGHT if c else PP_ALIGN.LEFT)
    for r, key in enumerate(ordered, start=1):
        items = groups[key]
        subtotal = sum(e["amount"] or 0 for e in items)
        for c, (txt, align) in enumerate((
            (group_label(key), PP_ALIGN.LEFT),
            (str(len(items)), PP_ALIGN.RIGHT),
            (inr(subtotal), PP_ALIGN.RIGHT),
        )):
            cell = table.cell(r, c)
            cell.text = txt
            _style_cell(cell, size=12, align=align)
    total_row = len(ordered) + 1
    for c, (txt, align) in enumerate((
        ("Total", PP_ALIGN.LEFT), ("", PP_ALIGN.RIGHT), (inr(grand), PP_ALIGN.RIGHT),
    )):
        cell = table.cell(total_row, c)
        cell.text = txt
        _style_cell(cell, size=13, bold=True, align=align, fill=RGBColor(0xE6, 0xE3, 0xDC))

    # ── Line items per shoot ───────────────────────────────────────────
    PER_SLIDE = 9
    for key in ordered:
        items = groups[key]
        subtotal = sum(e["amount"] or 0 for e in items)
        chunks = [items[i:i + PER_SLIDE] for i in range(0, len(items), PER_SLIDE)]
        for ci, chunk in enumerate(chunks):
            slide = _blank(prs)
            cont = f" (cont. {ci + 1}/{len(chunks)})" if len(chunks) > 1 else ""
            _header(slide, group_label(key) + cont,
                    f"{len(items)} item{'s' if len(items) != 1 else ''} · {inr(subtotal)}")
            n = len(chunk) + 1
            table = slide.shapes.add_table(
                n, 4, Inches(0.7), Inches(2.35), SLIDE_W - Inches(1.4),
                Inches(0.44 * n),
            ).table
            table.columns[0].width = Inches(3.1)
            table.columns[1].width = Inches(4.6)
            table.columns[2].width = Inches(2.1)
            table.columns[3].width = Inches(2.13)
            for c, label in enumerate(("Paid to", "What for", "Category", "Amount")):
                cell = table.cell(0, c)
                cell.text = label.upper()
                _style_cell(cell, size=10, bold=True, color=MUTED,
                            align=PP_ALIGN.RIGHT if c == 3 else PP_ALIGN.LEFT)
            for r, e in enumerate(chunk, start=1):
                flag = "" if e["amount"] is not None else "  (amount TBC)"
                vals = (
                    (e["vendor"] or "—", PP_ALIGN.LEFT, False),
                    ((e["description"] or e["filename"] or "—") + flag, PP_ALIGN.LEFT, False),
                    (CATEGORY_LABELS.get(e["category"] or "", e["category"] or "—"),
                     PP_ALIGN.LEFT, False),
                    (inr(e["amount"]) if e["amount"] is not None else "—",
                     PP_ALIGN.RIGHT, True),
                )
                for c, (txt, align, bold) in enumerate(vals):
                    cell = table.cell(r, c)
                    cell.text = txt
                    _style_cell(cell, size=11, align=align, bold=bold)

    # ── Receipts appendix ──────────────────────────────────────────────
    with_files = [e for e in expenses if e["has_receipt"]]
    if with_files:
        slide = _blank(prs)
        _text(slide, Inches(0.7), Inches(3.1), Inches(12), Inches(0.8),
              "Receipts", size=44, font=SERIF)
        _text(slide, Inches(0.7), Inches(4.05), Inches(12), Inches(0.5),
              f"{len(with_files)} receipt{'s' if len(with_files) != 1 else ''} attached",
              size=13, color=MUTED)
    for e in with_files:
        try:
            jpeg = (_receipt_jpeg(storage.get_bytes(e["s3_key"]), e.get("content_type"))
                    if e.get("s3_key") else None)
        except Exception:
            jpeg = None
        slide = _blank(prs)
        caption = " · ".join(x for x in (
            e["vendor"],
            inr(e["amount"]) if e["amount"] is not None else None,
            _fmt_day(e["date"]) if e["date"] else "General",
        ) if x)
        _text(slide, Inches(0.7), Inches(0.45), Inches(12), Inches(0.4),
              caption or e.get("filename") or "Receipt", size=13, bold=True)
        if jpeg:
            from PIL import Image
            img = Image.open(io.BytesIO(jpeg))
            iw, ih = img.size
            max_w, max_h = Inches(11.9), Inches(6.1)
            scale = min(max_w / iw, max_h / ih)
            w, h = int(iw * scale), int(ih * scale)
            slide.shapes.add_picture(
                io.BytesIO(jpeg),
                Emu(int((SLIDE_W - w) / 2)), Inches(1.05), Emu(w), Emu(h),
            )
        else:
            _text(slide, Inches(0.7), Inches(3.3), Inches(12), Inches(0.5),
                  f"Receipt file: {e.get('filename') or '(attached in Cue)'} — "
                  "open it from the Expenses tab", size=13, color=MUTED)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
