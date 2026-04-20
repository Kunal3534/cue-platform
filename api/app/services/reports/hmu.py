"""12 — Hair, Makeup & Prosthetics xlsx report."""
from openpyxl import Workbook

from ._common import (
    write_title_block, apply_print_settings, set_col_widths,
    FONT_HEADER_WHITE, FONT_BODY, FONT_BODY_BOLD,
    FILL_SECTION_HDR, CENTER, LEFT, BORDER,
)
from ..project_graph import ProjectGraph


def generate(pg: ProjectGraph, out_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "HMU"

    headers = ["CHARACTER", "DATE", "LOOK", "CATEGORY", "APPLY (MIN)", "VENDOR"]
    n = len(headers)
    write_title_block(ws, f"{pg.project.title} — HAIR, MAKEUP & PROSTHETICS", subtitle=pg.project.window or "", n_cols=n)

    for c, h in enumerate(headers, 1):
        cell = ws.cell(3, c)
        cell.value = h
        cell.font = FONT_HEADER_WHITE
        cell.fill = FILL_SECTION_HDR
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[3].height = 22

    r = 4
    for session in pg.hmu_sessions:
        row_vals = [
            session.character_name.title(),
            session.day_date.strftime("%-d %b %Y") if session.day_date else "",
            session.look_id or "",
            session.category.replace("_", " ").title() if session.category else "",
            session.apply_minutes or "",
            session.artist_vendor or "",
        ]
        for c, val in enumerate(row_vals, 1):
            cell = ws.cell(r, c)
            cell.value = val
            cell.font = FONT_BODY_BOLD if c == 1 else FONT_BODY
            cell.alignment = LEFT
            cell.border = BORDER
        r += 1

    set_col_widths(ws, {1: 24, 2: 14, 3: 10, 4: 22, 5: 12, 6: 28})
    ws.freeze_panes = "A4"
    apply_print_settings(ws)
    wb.save(out_path)
